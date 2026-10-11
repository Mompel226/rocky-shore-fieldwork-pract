#!/usr/bin/env python3
"""The recording sheet for the shore (downloads/fieldwork-recording-sheet.docx). Run from the repository folder.
The IB paper questions and their mark scheme are not built here: they are IB copyright, and the teacher hands them out."""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document()
s = d.styles["Normal"]; s.font.name = "Calibri"; s.font.size = Pt(11)
s.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for sec in d.sections:
    sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
def para(t, bold=False, size=None, space=4):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = bold
    if size: r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space); return p
def table(rows, widths, size=9):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths): gc.set(qn("w:w"), str(int(w * 567)))
    for i, row in enumerate(rows):
        tr = t.rows[i]._tr; trPr = tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if i: hgt = OxmlElement("w:trHeight"); hgt.set(qn("w:val"), "454"); hgt.set(qn("w:hRule"), "atLeast"); trPr.append(hgt)   # 0.8 cm to write in on the shore
        for j, v in enumerate(row):
            c = t.cell(i, j); c.width = Cm(widths[j]); c.text = ""
            r = c.paragraphs[0].add_run(v); r.font.size = Pt(size); r.bold = (i == 0)
    return t
para("Rocky shore fieldwork", bold=True, size=18)
para("Recording sheet for the shore", bold=True, size=14)
para("Transect: ........  Start time: ........  Tide: low tide at ........  Recorder: ........  Cover estimated by: ........", size=10)
rows = [["Distance\n/ m", "Time", "Rock temp.\n/ °C", "Barnacles\n/ %", "Green seaweed\n/ %", "Brown seaweed\n/ %", "Limpets\n/ number", "Other species\n/ % or number", "Notes"]]
rows += [[str(i), "", "", "", "", "", "", "", ""] for i in range(0, 21)]
table(rows, [1.6, 1.3, 1.6, 1.7, 1.8, 1.8, 1.6, 2.1, 2.7])
para("Estimate cover with the seaweed lifted for the animals under it. Covers can add up to more than 100 %. Photograph each quadrat.", size=9)
d.save("downloads/fieldwork-recording-sheet.docx")
