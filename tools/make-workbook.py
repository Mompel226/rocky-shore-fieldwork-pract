#!/usr/bin/env python3
"""The fieldwork workbook: opens in Google Sheets (upload, then Open with Google Sheets) and in Excel.
Yellow cells are for the student's data; every result is a formula they can read and rebuild."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.worksheet.datavalidation import DataValidation
fill_example = len(sys.argv) > 2 and sys.argv[2] == "example"
d = json.load(open("out2/data/shore.json")); S = d["stations"]; Q = [r for r in d["random"] if not r.get("skip")]
YEL = PatternFill("solid", fgColor="FFF4B8"); GRY = PatternFill("solid", fgColor="EEF0F5"); TEAL = PatternFill("solid", fgColor="DDF1F2")
B = Font(bold=True); H = Font(bold=True, size=14); MONO = Font(name="Courier New", size=10, color="555555")
thin = Side(style="thin", color="C9CCD6"); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
wb = Workbook()

# ---------- Start here
ws = wb.active; ws.title = "Start here"
lines = [("Rocky shore fieldwork: the workbook", H), ("IB Biology · B4.1.4 · C4.1.15 · https://nlcsbiology.com/rocky-shore-fieldwork-pract/", None), ("", None),
 ("1. In Google Drive: New ▸ File upload, choose this file, then right-click it ▸ Open with ▸ Google Sheets.", None),
 ("2. Type your data in the YELLOW cells only: the Transect tab and the Random quadrats tab.", None),
 ("3. Every other number is a formula. Click a cell to read its formula in the formula bar; the grey text beside it shows the same formula.", None),
 ("4. The Example data tab holds the virtual shore's data, so you can check that your formulas give the same answers.", None), ("", None),
 ("Make the scatter graph yourself (Google Sheets):", B),
 ("   Select the temperature column and one species column (hold Ctrl or ⌘ to select two columns that are not next to each other).", None),
 ("   Insert ▸ Chart. Chart editor ▸ Setup ▸ Chart type ▸ Scatter chart.", None),
 ("   Customise ▸ Series: tick Trendline; Label ▸ Use equation; tick Show R².", None),
 ("   Give the graph a title, for example: Figure 1. The effect of rock temperature on barnacle cover.", None), ("", None),
 ("Code: AGPL-3.0 · Teaching material: CC BY-NC-SA 4.0 · Dr Daniel Mompel Riera, NLCS Jeju", Font(size=9, color="666666"))]
for i, (t, f) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=t)
    if f: c.font = f
ws.column_dimensions["A"].width = 120

# ---------- Transect
wt = wb.create_sheet("Transect")
head = ["Distance / m", "Rock temperature / °C", "Barnacles / %", "Sea lettuce / %", "Brown seaweed / %", "Limpets / number"]
for j, hd in enumerate(head, 1):
    c = wt.cell(row=1, column=j, value=hd); c.font = B; c.fill = GRY; c.alignment = Alignment(wrap_text=True, horizontal="center"); c.border = BOX
for i, s in enumerate(S):
    r = i + 2
    wt.cell(row=r, column=1, value=s["d"]).border = BOX
    c = wt.cell(row=r, column=2, value=s["temp"] if fill_example else None); c.fill = YEL; c.border = BOX
    for j, k in enumerate(["barn", "green", "brown", "limp"], 3):
        c = wt.cell(row=r, column=j, value=s[k] if fill_example else None); c.fill = YEL; c.border = BOX
for j in range(1, 7): wt.column_dimensions["ABCDEFGH"[j - 1]].width = 15
wt.column_dimensions["G"].width = 3; wt.column_dimensions["H"].width = 40; wt.column_dimensions["I"].width = 46
wt.cell(row=1, column=8, value="Correlation with rock temperature").font = B
rows = [("Barnacles", "C"), ("Sea lettuce", "D"), ("Brown seaweed", "E"), ("Limpets", "F")]
r0 = 3
for n, (nm, col) in enumerate(rows):
    r = r0 + n * 6
    wt.cell(row=r, column=8, value=nm).font = B
    items = [("Pearson's r", f'=IFERROR(CORREL($B$2:$B$12,{col}2:{col}12),"")', f"=CORREL(B2:B12, {col}2:{col}12)"),
             ("R² (straight line)", f'=IFERROR(RSQ({col}2:{col}12,$B$2:$B$12),"")', f"=RSQ({col}2:{col}12, B2:B12)"),
             ("Number of quadrats", f'=COUNT({col}2:{col}12)', f"=COUNT({col}2:{col}12)"),
             ("Critical r (p = 0.05)", f"=IFERROR(VLOOKUP(I{r+3},'Critical values'!$A$4:$B$29,2,FALSE),\"\")", "looks up the number of quadrats in Critical values"),
             ("Significant?", f'=IF(I{r+1}="","",IF(ABS(I{r+1})>I{r+4},"yes: significant (p < 0.05)","no: not significant"))', "=IF(ABS(r) > critical r, …)")]
    for m, (lab, f, show) in enumerate(items):
        wt.cell(row=r + 1 + m, column=8, value="   " + lab)
        c = wt.cell(row=r + 1 + m, column=9, value=f); c.fill = TEAL; c.border = BOX
        if m < 2: c.number_format = "0.00"
        c2 = wt.cell(row=r + 1 + m, column=10, value=show); c2.data_type = 's'; c2.font = MONO
wt.column_dimensions["J"].width = 44
wt.cell(row=28, column=8, value="Look at the scatter graph before you trust r: r measures only a straight-line pattern.").font = Font(italic=True, color="8F5310")
ch = ScatterChart(); ch.title = "Barnacle cover and rock temperature"; ch.style = 13
ch.x_axis.title = "Rock temperature / °C"; ch.y_axis.title = "Barnacle cover / %"
xs = Reference(wt, min_col=2, min_row=2, max_row=12); ys = Reference(wt, min_col=3, min_row=2, max_row=12)
se = Series(ys, xs, title="Barnacles"); se.marker.symbol = "circle"; se.graphicalProperties.line.noFill = True
from openpyxl.chart.trendline import Trendline
se.trendline = Trendline(trendlineType="linear", dispRSqr=True, dispEq=True); ch.series.append(se)
ch.x_axis.scaling.min = 22; ch.x_axis.scaling.max = 36; ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100
ch.height = 8; ch.width = 14
wt.add_chart(ch, "A15")

# ---------- Random quadrats
wr = wb.create_sheet("Random quadrats")
for j, hd in enumerate(["Quadrat", "Across (0–9)", "Down (0–9)", "Limpets", "Sea lettuce"], 1):
    c = wr.cell(row=1, column=j, value=hd); c.font = B; c.fill = GRY; c.border = BOX
dv = DataValidation(type="list", formula1='"present,absent"', allow_blank=True); wr.add_data_validation(dv)
for i in range(30):
    r = i + 2; wr.cell(row=r, column=1, value=i + 1).border = BOX
    q = Q[i]
    for j, v in ((2, q["gx"]), (3, q["gy"]), (4, "present" if q["limp"] else "absent"), (5, "present" if q["green"] else "absent")):
        c = wr.cell(row=r, column=j, value=v if fill_example else None); c.fill = YEL; c.border = BOX
        if j >= 4: dv.add(c)
for col, w in zip("ABCDEFG", (10, 13, 12, 12, 13, 3, 58)): wr.column_dimensions[col].width = w
wr.cell(row=1, column=7, value="Random positions").font = B
wr.cell(row=2, column=7, value="=RANDBETWEEN(0,9)")
wr.cell(row=2, column=8, value="=RANDBETWEEN(0,9)")
wr.cell(row=3, column=7, value="The cell above gives a new random number from 0 to 9 every time the sheet changes.").alignment = wrap
wr.cell(row=4, column=7, value="Write each pair of numbers down as soon as you get it: they change.").alignment = wrap

# ---------- Chi-squared
wc = wb.create_sheet("Chi-squared")
wc["A1"] = "Chi-squared test for association between limpets and sea lettuce"; wc["A1"].font = H
wc["A2"] = "H0: there is no association between the distribution of limpets and the distribution of sea lettuce."
wc["A3"] = "Observed (O): counted from the Random quadrats tab"; wc["A3"].font = B
lay = [("", "Sea lettuce present", "Sea lettuce absent", "Total")]
for j, v in enumerate(lay[0], 1):
    c = wc.cell(row=4, column=j, value=v); c.font = B; c.fill = GRY; c.border = BOX
wc["A5"] = "Limpets present"; wc["A6"] = "Limpets absent"; wc["A7"] = "Total"
RQ = "'Random quadrats'!"
wc["B5"] = f'=COUNTIFS({RQ}$D$2:$D$31,"present",{RQ}$E$2:$E$31,"present")'
wc["C5"] = f'=COUNTIFS({RQ}$D$2:$D$31,"present",{RQ}$E$2:$E$31,"absent")'
wc["B6"] = f'=COUNTIFS({RQ}$D$2:$D$31,"absent",{RQ}$E$2:$E$31,"present")'
wc["C6"] = f'=COUNTIFS({RQ}$D$2:$D$31,"absent",{RQ}$E$2:$E$31,"absent")'
for r in (5, 6): wc[f"D{r}"] = f"=SUM(B{r}:C{r})"
for col in "BCD": wc[f"{col}7"] = f"=SUM({col}5:{col}6)"
wc["A9"] = "Expected (E) = row total × column total ÷ grand total"; wc["A9"].font = B
for j, v in enumerate(["", "Sea lettuce present", "Sea lettuce absent"], 1):
    c = wc.cell(row=10, column=j, value=v); c.font = B; c.fill = GRY; c.border = BOX
wc["A11"] = "Limpets present"; wc["A12"] = "Limpets absent"
wc["B11"] = '=IFERROR($D5*B$7/$D$7,"")'; wc["C11"] = '=IFERROR($D5*C$7/$D$7,"")'
wc["B12"] = '=IFERROR($D6*B$7/$D$7,"")'; wc["C12"] = '=IFERROR($D6*C$7/$D$7,"")'
wc["A14"] = "(O − E)² ÷ E"; wc["A14"].font = B
for j, v in enumerate(["", "Sea lettuce present", "Sea lettuce absent"], 1):
    c = wc.cell(row=15, column=j, value=v); c.font = B; c.fill = GRY; c.border = BOX
wc["A16"] = "Limpets present"; wc["A17"] = "Limpets absent"
for r, ro in ((16, 5), (17, 6)):
    for col in "BC": wc[f"{col}{r}"] = f'=IFERROR(({col}{ro}-{col}{ro + 6})^2/{col}{ro + 6},"")'
res = [("χ² (the sum of the four cells)", "=IFERROR(SUM(B16:C17),\"\")", "=SUM(B16:C17)"),
       ("Degrees of freedom", "=(2-1)*(2-1)", "(rows − 1) × (columns − 1)"),
       ("Critical value at p = 0.05", "=_xlfn.CHISQ.INV.RT(0.05,B20)", "=CHISQ.INV.RT(0.05, 1)"),
       ("p-value", "=IFERROR(_xlfn.CHISQ.TEST(B5:C6,B11:C12),\"\")", "=CHISQ.TEST(B5:C6, B11:C12)"),
       ("Smallest expected value (must be at least 5)", "=IFERROR(MIN(B11:C12),\"\")", "=MIN(B11:C12)"),
       ("Conclusion", '=IF(B19="","",IF(B19>B21,"χ² > critical value: reject H0. There is a significant association (p < 0.05).","χ² < critical value: do not reject H0. No significant association."))', "")]
for i, (lab, f, show) in enumerate(res):
    r = 19 + i; wc[f"A{r}"] = lab; c = wc[f"B{r}"]; c.value = f; c.fill = TEAL; c.border = BOX
    if i in (0, 2, 4): c.number_format = "0.00"
    if i == 3: c.number_format = "0.0000"
    wc[f"D{r}"] = show; wc[f"D{r}"].data_type = 's'; wc[f"D{r}"].font = MONO
wc["A26"] = "Then compare O with E: if the species are found together less often than expected, the association is negative."
wc["A27"] = "An association does not show the cause. It may come from grazing, from different niches or from competition."
wc.column_dimensions["A"].width = 44; wc.column_dimensions["B"].width = 20; wc.column_dimensions["C"].width = 20; wc.column_dimensions["D"].width = 34

# ---------- Example data (the virtual shore)
we = wb.create_sheet("Example data")
we["A1"] = "The virtual shore on the page (simulated data). Copy these into the yellow cells to check your formulas."; we["A1"].font = B
for j, hd in enumerate(head, 1): we.cell(row=3, column=j, value=hd).font = B
for i, s in enumerate(S): 
    for j, v in enumerate([s["d"], s["temp"], s["barn"], s["green"], s["brown"], s["limp"]], 1): we.cell(row=4 + i, column=j, value=v)
for j, hd in enumerate(["Quadrat", "Across", "Down", "Limpets", "Sea lettuce"], 8): we.cell(row=3, column=j, value=hd).font = B
for i, q in enumerate(Q):
    for j, v in enumerate([i + 1, q["gx"], q["gy"], "present" if q["limp"] else "absent", "present" if q["green"] else "absent"], 8): we.cell(row=4 + i, column=j, value=v)
for col in "ABCDEFHIJKL": we.column_dimensions[col].width = 14

# ---------- Critical values
wv = wb.create_sheet("Critical values")
wv["A1"] = "Critical values of Pearson's r, p = 0.05, two-tailed"; wv["A1"].font = B
wv["A3"] = "Number of pairs (n)"; wv["B3"] = "Critical r"; wv["A3"].font = B; wv["B3"].font = B
CR = {5: .878, 6: .811, 7: .754, 8: .707, 9: .666, 10: .632, 11: .602, 12: .576, 13: .553, 14: .532, 15: .514, 16: .497, 17: .482, 18: .468, 19: .456, 20: .444,
      21: .433, 22: .423, 23: .413, 24: .404, 25: .396, 26: .388, 27: .381, 28: .374, 29: .367, 30: .361}
for i, (n, v) in enumerate(CR.items()): wv.cell(row=4 + i, column=1, value=n); wv.cell(row=4 + i, column=2, value=v)
wv["D1"] = "Critical values of χ²"; wv["D1"].font = B
wv["D3"] = "Degrees of freedom"; wv["E3"] = "p = 0.05"; wv["F3"] = "p = 0.01"
for c in ("D3", "E3", "F3"): wv[c].font = B
for i, (df, a, b) in enumerate([(1, 3.84, 6.63), (2, 5.99, 9.21), (3, 7.81, 11.34), (4, 9.49, 13.28)]):
    wv.cell(row=4 + i, column=4, value=df); wv.cell(row=4 + i, column=5, value=a); wv.cell(row=4 + i, column=6, value=b)
wv.column_dimensions["A"].width = 20; wv.column_dimensions["D"].width = 20
wb.save(sys.argv[1]); print("saved", sys.argv[1])
